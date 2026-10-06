from pathlib import Path
import re
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "outputs" / "本地Python_AI助手新手说明书.md"
TEMPLATE = ROOT / "科研研究提示词模板.md"
DESKTOP = Path.home() / "Desktop"
OUT = DESKTOP / "本地Python_AI助手新手说明书_中文版.docx"


def set_cell_shading(paragraph, fill="F3F4F6"):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)


def add_text(doc, text, bold=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    set_run_font(r, bold=bold)
    return p


def add_markdown(doc, text):
    code = False
    for raw in text.replace("\r\n", "\n").split("\n"):
        line = raw.rstrip()
        if line.strip().startswith("```"):
            code = not code
            continue
        if code:
            p = doc.add_paragraph(style="No Spacing")
            set_cell_shading(p, "F3F4F6")
            p.paragraph_format.left_indent = Inches(0.25)
            r = p.add_run(line)
            set_run_font(r, name="Consolas", size=9)
            continue
        if not line.strip():
            continue
        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            for r in p.runs: set_run_font(r, size=18, bold=True, color=(31, 78, 121))
        elif line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            for r in p.runs: set_run_font(r, size=14, bold=True, color=(31, 78, 121))
        elif line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            for r in p.runs: set_run_font(r, size=12, bold=True, color=(68, 68, 68))
        elif re.match(r"^\s*(?:[-*]|\d+\.)\s+", line):
            item = re.sub(r"^\s*(?:[-*]|\d+\.)\s+", "", line)
            p = doc.add_paragraph(style="List Bullet" if not re.match(r"^\s*\d+\.", line) else "List Number")
            r = p.add_run(item)
            set_run_font(r)
        elif line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.35)
            r = p.add_run(line[2:])
            set_run_font(r, color=(89, 89, 89))
            set_cell_shading(p, "EAF2F8")
        else:
            p = doc.add_paragraph()
            # Basic inline formatting while retaining readable URLs and commands.
            pieces = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", line)
            for piece in pieces:
                if not piece: continue
                if piece.startswith("`") and piece.endswith("`"):
                    r = p.add_run(piece[1:-1]); set_run_font(r, name="Consolas", size=9.5, color=(128, 0, 0))
                elif piece.startswith("**") and piece.endswith("**"):
                    r = p.add_run(piece[2:-2]); set_run_font(r, bold=True)
                else:
                    r = p.add_run(piece); set_run_font(r)


def main():
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Inches(0.7)
    sec.bottom_margin = Inches(0.7)
    sec.left_margin = Inches(0.8)
    sec.right_margin = Inches(0.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)

    add_markdown(doc, SOURCE.read_text(encoding="utf-8"))
    doc.add_page_break()
    add_markdown(doc, TEMPLATE.read_text(encoding="utf-8"))

    footer = doc.sections[0].footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("本地 Python AI 助手 · 离线使用说明")
    set_run_font(run, size=9, color=(128, 128, 128))
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
